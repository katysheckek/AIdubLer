
############### превратить разные форматы файлов в единую структуру данных.

"""
независимо от того, был ли исходник PDF или презентацией,
дальше мы хотим работать с объектами примерно такого вида:
{
    "text": "Фотосинтез — это процесс...",
    "source": "biology.pdf",
    "location": "page:12"
}
"""

'''
PDF сохраняет номер страницы.
Презентация сохраняет номер слайда.
DOCX пока читается как единый текст. 
Номер страницы в DOCX не фиксирован: он зависит от форматирования и программы, в которой документ открывается.
>>>    Сканированные PDF могут не содержать извлекаемого текста. Для них позднее потребуется OCR.
Изображения, схемы, графики и заметки презентаций пока не обрабатываются полноценно.
'''

from pathlib import Path

import pymupdf
from docx import Document
from pptx import Presentation


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx"}


def load_pdf(path: Path) -> list[dict]:
    result = []

    with pymupdf.open(path) as pdf:
        for page_number, page in enumerate(pdf, start=1):
            text = page.get_text("text").strip()

            if text:
                result.append({
                    "text": text,
                    "source": path.name,
                    "location": f"page:{page_number}",
                })

    return result


def load_docx(path: Path) -> list[dict]:
    doc = Document(path)
    paragraphs = [
        p.text.strip()
        for p in doc.paragraphs
        if p.text.strip()
    ]

    text = "\n".join(paragraphs)

    if not text:
        return []

    return [{
        "text": text,
        "source": path.name,
        "location": "document",
    }]


def load_pptx(path: Path) -> list[dict]:
    presentation = Presentation(path)
    result = []

    for slide_number, slide in enumerate(
        presentation.slides, start=1
    ):
        parts = []

        for shape in slide.shapes:
            if shape.has_text_frame and shape.text.strip():
                parts.append(shape.text.strip())

            if shape.has_table:
                for row in shape.table.rows:
                    parts.append(
                        " | ".join(
                            cell.text.strip()
                            for cell in row.cells
                        )
                    )

        text = "\n".join(parts).strip()

        if text:
            result.append({
                "text": text,
                "source": path.name,
                "location": f"slide:{slide_number}",
            })

    return result


def load_document(path: Path) -> list[dict]:
    extension = path.suffix.lower()

    if extension == ".pdf":
        return load_pdf(path)

    if extension == ".docx":
        return load_docx(path)

    if extension == ".pptx":
        return load_pptx(path)

    raise ValueError(f"Unsupported format: {extension}")


def load_folder(folder: str) -> list[dict]:
    documents = []
    root = Path(folder)

    for path in sorted(root.rglob("*")):
        if (
            path.is_file()
            and path.suffix.lower() in SUPPORTED_EXTENSIONS
        ):
            documents.extend(load_document(path))

    return documents